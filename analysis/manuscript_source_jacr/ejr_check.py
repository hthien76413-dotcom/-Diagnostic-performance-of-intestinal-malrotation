# -*- coding: utf-8 -*-
"""Pre-submission checks of the built EJR files against the journal's Guide for
Authors (printed 2026-10-15; analysis/manuscript_source_jacr/EJR_Guide_for_Authors_2026-10-15.pdf).

Run after the builders. Every check prints PASS or FAIL (INFO lines are
reported, not scored); the script exits non-zero if any check fails.
"""
import datetime, docx, glob, io, os, re, subprocess, sys, zipfile
from docx.oxml.ns import qn
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from usspell import british_left

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = '/home/user/-Diagnostic-performance-of-intestinal-malrotation/'
fails = 0


def check(ok, msg):
    global fails
    print(('PASS ' if ok else 'FAIL ') + msg)
    fails += (not ok)


def info(msg):
    print('INFO ' + msg)


def paras(path):
    d = docx.Document(path)
    out = []
    for el in d.element.body.iterchildren():
        if el.tag in (qn('w:p'), qn('w:tbl')):
            t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            if t.strip():
                out.append(t)
    return d, out


F = {k: ROOT + v for k, v in dict(
    cover='EJR_1_CoverLetter.docx', title='EJR_2_TitlePage.docx', ms='EJR_3_Manuscript_anonymized.docx',
    hl='EJR_4_Highlights.docx', decl='EJR_5_Declarations.docx', credit='EJR_6_CreditAuthorStatement.docx',
    strobe='EJR_STROBE_checklist.docx', s1='EJR_Supplement_1_NLP_and_report_audit.docx',
    s2='EJR_Supplement_2_models_and_subgroups.docx', s3='EJR_Supplement_3_CT_and_UGI_content_audit.docx',
    s4='EJR_Supplement_4_minimum_reporting_dataset.docx').items()}
for k, p in F.items():
    check(os.path.exists(p), f'{os.path.basename(p)} exists')
ms_doc, ms = paras(F['ms'])
_, tp = paras(F['title'])
_, hl = paras(F['hl'])
_, dc = paras(F['decl'])
_, cr = paras(F['credit'])
_, cv = paras(F['cover'])

# ---------------------------------------------------------------- title, abstract, keywords
TITLE = ms[0]
check(TITLE == tp[0], 'title identical on the title page and in the manuscript')
check(not re.search(r'\b[A-Z]{2,}\b', TITLE) and not TITLE.rstrip().endswith('?'), f'title has no abbreviations ({len(TITLE)} characters)')
i = ms.index('Abstract')
abstract = ms[i + 1:i + 5]
words = sum(len(p.split()) for p in abstract)
check([p.split(':')[0] for p in abstract] == ['Objectives', 'Methods', 'Results', 'Conclusions'], 'abstract paragraphs Objectives/Methods/Results/Conclusions')
check(words <= 250, f'abstract {words} words incl. headings (limit 250)')
check(not any(re.search(r'\[\d', p) for p in abstract), 'no reference citations in the abstract')
abbr = sorted({m for p in abstract for m in re.findall(r'\b[A-Z]{2,}\b', p)})
check(not abbr, f'no abbreviations in the abstract {abbr if abbr else ""}')
kw = [p for p in ms if p.startswith('Keywords:')][0].split(':', 1)[1].split(';')
check(1 <= len(kw) <= 7 and not any(re.search(r'\b(and|of)\b', k) for k in kw), f'{len(kw)} keywords (1-7), none with "and"/"of"')

# ---------------------------------------------------------------- highlights
bul = [p.lstrip('• ').strip() for p in hl if p.startswith('•')]
check(3 <= len(bul) <= 5 and all(len(b) <= 85 for b in bul), f'{len(bul)} highlights, longest {max(map(len, bul))} characters (3-5 bullets, max 85)')
check('highlights' in os.path.basename(F['hl']).lower(), 'highlights file name contains the word "highlights"')

# ---------------------------------------------------------------- structure of the anonymized manuscript
heads = [p.style.name for p in ms_doc.paragraphs]
h1 = [p.text for p in ms_doc.paragraphs if p.style.name == 'Heading 1']
want = ['Abstract', 'Introduction', 'Materials and Methods', 'Results', 'Discussion', 'Conclusion']
check([h for h in h1 if h in want] == want, 'sections Abstract, Introduction, Materials and Methods, Results, Discussion, Conclusion in order')
AI_HEAD = 'Declaration of generative AI and AI-assisted technologies in the manuscript preparation process'
check(AI_HEAD in h1 and h1.index(AI_HEAD) < h1.index('References') and h1.index(AI_HEAD) == h1.index('References') - 1,
      'AI declaration in its own section with the journal\'s title, directly before the reference list')
ai_txt = ms[ms.index(AI_HEAD) + 1]
check(ai_txt.startswith('During the preparation of this work the authors used Claude (Anthropic) to assist with drafting and language editing of the manuscript and with writing and checking the analysis code.')
      and 'take full responsibility for the content of the published article' in ai_txt, 'AI declaration names the tool and the full scope of use (drafting, language editing, analysis code)')
check(any(p == 'Limitations' for p in [p.text for p in ms_doc.paragraphs if p.style.name == 'Heading 2']), 'Limitations subsection')
check(not any(re.match(r'(Acknowledg|Funding|Author contributions|CRediT)', p) for p in ms), 'no acknowledgments, funding or contributions in the anonymized manuscript')
check(len(ms_doc.inline_shapes) == 0, 'no figures embedded in the manuscript (figures are separate files)')
check(h1.index('Figure Legends') > h1.index('References'), 'figure legends follow the references')

# ---------------------------------------------------------------- word counts (EJR: 4,000 words incl. abstract, legends and up to 40 references)
body = int(subprocess.check_output([sys.executable, os.path.join(HERE, 'wc.py')]).strip())
leg = [p for p in ms[ms.index('Figure Legends') + 1:] if re.match(r'Figure \d\. ', p)]
legw = sum(len(p.split()) for p in leg)
refs = [p for p in ms if re.match(r'\[\d+\] ', p)]
refw = sum(len(p.split()) for p in refs)
check(len(leg) == 3, '3 figure legends')
check(body + words + legw <= 4000, f'main text {body} + abstract {words} + legends {legw} = {body + words + legw} words (limit 4,000)')
info(f'with the {refw} reference words the total is {body + words + legw + refw}; the guide says "including abstract, legends and up to 40 references", which can be read as counting references')

# ---------------------------------------------------------------- tables
tabs = ms_doc.tables
check(len(tabs) == 4, f'{len(tabs)} tables')
bad = []
for k, t in enumerate(tabs, 1):
    x = t._tbl.xml
    if re.search(r'<w:(left|right|insideV)\b[^>]*w:val="(?!nil)', x) or '<w:shd ' in x:
        bad.append(k)
check(not bad, f'no vertical rules or shading in the tables {bad if bad else ""}')
body_xml = ms_doc.element.body
prev_break = []
els = list(body_xml.iterchildren())
for j, el in enumerate(els):
    if el.tag == qn('w:p') and re.match(r'Table \d\. ', ''.join(x.text or '' for x in el.iter(qn('w:t')))):
        prev_break.append(any('w:br' in e.xml and 'type="page"' in e.xml for e in els[max(0, j - 1):j + 1] if e.tag == qn('w:p')))
check(len(prev_break) == 4 and all(prev_break), 'every table starts on its own page')
t1 = ' '.join(p for p in ms if p.startswith('Groups overlap'))
check('Sex is as recorded in the medical record' in t1, 'Table 1 note defines sex (SAGER guidance)')

# ---------------------------------------------------------------- references
cited = set()
for p in ms[:ms.index('References')]:
    for grp in re.findall(r'\[([\d,–-]+)\]', p):
        for part in grp.split(','):
            lo, _, hi = part.replace('-', '–').partition('–')
            cited |= set(range(int(lo), int(hi or lo) + 1))
check(len(refs) <= 40 and cited == set(range(1, len(refs) + 1)), f'{len(refs)} references (max 40), every one cited, none missing')
check([int(re.match(r'\[(\d+)\]', r).group(1)) for r in refs] == list(range(1, len(refs) + 1)), 'references numbered consecutively')
first = []
for p in ms[:ms.index('References')]:
    for grp in re.findall(r'\[([\d,–-]+)\]', p):
        for part in grp.split(','):
            lo, _, hi = part.replace('-', '–').partition('–')
            for n in range(int(lo), int(hi or lo) + 1):
                if n not in first:
                    first.append(n)
check(first == sorted(first), 'references first cited in numerical order')
check(not any(re.search(r'\[\d+, \d+', p) for p in ms), 'citations written without a space after the comma, e.g. [3,6]')
check(all(re.search(r'\(\d{4}\) \S+( – \S+)?\. https://doi\.org/10\.', r) for r in refs), 'references in the Elsevier numeric layout with a DOI link')

# ---------------------------------------------------------------- anonymization and identity
IDENT = re.compile(r'Wuhan|Tongji|Huazhong|2026R018|Jun Yang|Jun Shu|Bian\b|Zhengliang|Haiyan|Kai Zheng|Fei Peng|yjun|Guanghua|Hongxi|Haibin|orcid|Hong Kong Road')
for k in ('ms', 'decl', 'strobe', 's1', 's2', 's3', 's4', 'hl'):
    d, ps = paras(F[k])
    hits = sorted({m.group(0) for p in ps for m in IDENT.finditer(p)})
    check(not hits, f'{os.path.basename(F[k])}: no identifying text {hits if hits else ""}')
tp_all = ' '.join(tp)
for name in ('Guanghua Zhang', 'Hongxi Guo', 'Haibin Wang'):
    where = [os.path.basename(F[k]) for k in F if any(name in p for p in paras(F[k])[1])]
    check(where == ['EJR_2_TitlePage.docx'], f'{name} appears only on the title page')

# ---------------------------------------------------------------- title page
authors = ['Jun Shu', 'Fei Peng', 'Kai Zheng', 'Haiyan Lei', 'Zhengliang Meng', 'Hongqiang Bian', 'Jun Yang']
auth_line = [p for p in tp if p.startswith('Authors:')][0]
check([a for a in authors if a in auth_line] == authors and auth_line.index('Jun Shu') < auth_line.index('Fei Peng') < auth_line.index('Kai Zheng') < auth_line.index('Haiyan Lei') < auth_line.index('Zhengliang Meng') < auth_line.index('Hongqiang Bian') < auth_line.index('Jun Yang'), '7 authors on the title page in the submission order, given name before family name')
tp_doc = docx.Document(F['title'])
sup = [r.text for p in tp_doc.paragraphs for r in p.runs if r.font.superscript]
check({'a', 'b'} <= set(s.strip('*,') for s in sup[:20]) or any(s.startswith('a') for s in sup), 'affiliation marks are superscript letters')
aff = [p for p in tp if p.startswith('Affiliations:')][0]
check(aff.count('China') == 2 and aff.count('Wuhan 430016') == 2 and aff.count('Hong Kong Road') == 2, 'both affiliations carry a full postal address with the country')
corr = [p for p in tp if p.startswith('*Corresponding author:')][0]
check('yjun201602@163.com' in corr and 'Hong Kong Road' in corr and 'China' in corr, 'corresponding author: full address and email')
check('[APPROVAL DATE]' not in tp_all, 'ethics approval date filled in on the title page (placeholder [APPROVAL DATE] still present)' if '[APPROVAL DATE]' in tp_all else 'ethics approval date filled in on the title page')
for lead in ('Acknowledgments:', 'Funding:', 'Declaration of competing interests:', 'Ethics approval and consent to participate:', 'Consent for publication:'):
    check(any(p.startswith(lead) for p in tp), f'title page has "{lead}"')
check(any('This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.' in p for p in tp), 'funding sentence is the journal\'s standard wording')
check(any('2026R018-E01' in p for p in tp), 'ethics reference number on the title page')

# ---------------------------------------------------------------- CRediT file
check([p.split(':')[0] for p in cr[1:]] == authors, 'CRediT file lists the 7 authors, name first')
roles = {'Conceptualization', 'Data curation', 'Formal analysis', 'Funding acquisition', 'Investigation', 'Methodology', 'Project administration', 'Resources', 'Software', 'Supervision', 'Validation', 'Visualization', 'Writing – original draft', 'Writing – review & editing'}
used = {r.strip().rstrip('.') for p in cr[1:] for r in p.split(':', 1)[1].split(',')}
check(used <= roles, f'only CRediT roles are used {sorted(used - roles) if used - roles else ""}')

# ---------------------------------------------------------------- supplements cited, spelling, stale terms
text_ms = '\n'.join(ms[:ms.index('References')])
for n in (1, 2, 3, 4):
    check(f'Supplement {n}' in text_ms, f'Supplement {n} is cited in the manuscript text')
stale = re.compile(r'JACR|\bJUM\b|Journal of the American College of Radiology|Journal of Ultrasound in Medicine|Take-Home|Practitioner Points|Online Resource|Insights into Imaging|Graphical abstract|\bFig\. |171/320|65/119|54\.6%|237/301|169/320|64/119|58 of 113|58/113|\b410 children')
stale_post = re.compile(r'\b740 index|\b119 ultrasound|of the 119\b')
for k in F:
    _, ps = paras(F[k])
    t = '\n'.join(p for p in ps if not re.match(r'\[\d+\] ', p))
    check(not british_left(t), f'{os.path.basename(F[k])}: American spelling {british_left(t) or ""}')
    hits = sorted({m.group(0) for m in stale.finditer(t)})
    if k != 's1':
        hits = sorted(set(hits) | {m.group(0) for m in stale_post.finditer(t)})
    check(not hits, f'{os.path.basename(F[k])}: no stale terms or retired values {hits if hits else ""}')
check(any('European Journal of Radiology' in p for p in cv) and any('Article Transfer Service' in p for p in cv), 'cover letter is addressed to EJR and states the transfer')

# ---------------------------------------------------------------- operation sheet carries the abstract verbatim
ops = '\n'.join(paras(ROOT + 'EJR_投稿操作单.docx')[1])
check(all(p in ops for p in abstract), 'operation sheet abstract equals the manuscript abstract')

# ---------------------------------------------------------------- figures: vector PDF with embedded fonts, no producer tags
for k in (1, 2, 3):
    b = open(ROOT + f'Figure_{k}.pdf', 'rb').read()
    check(b.startswith(b'%PDF') and b'/FontFile' in b and not re.search(rb'/(Producer|Creator|CreationDate|Author|Title)\b', b), f'Figure_{k}.pdf: vector, fonts embedded, no producer or date tags')

# ---------------------------------------------------------------- file properties of every docx
for f in sorted(glob.glob(ROOT + 'EJR_*.docx')):
    z = zipfile.ZipFile(f)
    core = z.read('docProps/core.xml').decode('utf-8')
    empty = all(re.search(rf'<{t}></{t}>|<{t}/>', core) for t in ('dc:creator', 'cp:lastModifiedBy', 'dc:description', 'dc:title'))
    years = set(re.findall(r'<dcterms:(?:created|modified)[^>]*>(\d{4})-', core))
    parts = z.namelist()
    pngs = [n for n in parts if n.startswith('word/media/') and n.lower().endswith('.png')]
    clean_png = all(not [k for k in Image.open(io.BytesIO(z.read(n))).info if k != 'dpi'] for n in pngs)
    check(empty and years == {str(datetime.date.today().year)} and 'docProps/thumbnail.jpeg' not in parts and clean_png,
          f'{os.path.basename(f)}: properties empty, dated this year, no template thumbnail')

print('\n%d check(s) failed' % fails if fails else '\nall checks passed')
sys.exit(1 if fails else 0)
