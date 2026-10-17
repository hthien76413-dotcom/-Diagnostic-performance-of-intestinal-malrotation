import docx, json, re, os
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from threeline import make_three_line_table
from usspell import us, us_table, ama_p
from docstyle import apply_house_style, add_page_numbers
import shutil

def _blank_props(doc):
    """Strip python-docx's default authorship before saving.

    These files go out for double-blind review, so no document property may
    name an author, a reviser or the generator.
    """
    c = doc.core_properties
    for f in ('author', 'last_modified_by', 'title', 'subject', 'comments',
              'category', 'keywords', 'content_status', 'identifier',
              'language', 'version'):
        setattr(c, f, '')
    return doc


AN='/home/user/-Diagnostic-performance-of-intestinal-malrotation/analysis/'
OUT='/home/user/-Diagnostic-performance-of-intestinal-malrotation/'
D=json.load(open(AN+'tables123.json')); D.update(json.load(open(AN+'tables456.json')))
D.update(json.load(open(AN+'or_tables.json'))); D.update(json.load(open(AN+'or_h.json')))
D.update(json.load(open(AN+'or3_pooled.json'))); D.update(json.load(open(AN+'or_sens.json'))); D.update(json.load(open(AN+'or_add.json')))
D.update(json.load(open(AN+'us_coding_agreement.json'))); D.update(json.load(open(AN+'mrd.json'))); D.update(json.load(open(AN+'hp_tables.json')))
# Table S1 in English: the reports are Chinese, so the third column gives the English translation of the terms each
# pattern matches; the original regular expressions are appended verbatim to Supplement_1_classifier.py (CONTENT_PATTERNS).
H_EN = {
 ('Third portion of duodenum (D3)', 'Ultrasound'): '"horizontal portion", "horizontal segment", "transverse portion" or "third segment" of the duodenum; or "transverse portion" or "horizontal portion" on its own, except "transverse portion" directly preceded by the character for "arch" (transverse portion of the aortic arch)',
 ('Duodenojejunal junction', 'Ultrasound'): '"duodenojejunal flexure"; the Chinese short form of the ligament of Treitz; "Treitz"; "suspensory ligament of the duodenum"; "duodenum-jejunum junction" (with an optional hyphen, "with" or "and" between the two words)',
 ('Duodenum mentioned in any form', 'Ultrasound'): '"duodenum" (any mention)',
 ('Superior mesenteric artery–vein relationship', 'Ultrasound / CT'): '"superior mesenteric artery and vein" (with or without a separating enumeration comma); "superior mesenteric artery"; "superior mesenteric vein"; "mesenteric vessels"; SMA; SMV; excluding matches whose only context is the left renal vein, nutcracker phenomenon or renal-vein compression with no whirlpool term',
 ('Explicit vessel inversion', 'Ultrasound'): '"vein" or "artery" followed within 10 characters in the same clause by "transposition", "reversed position", "inversion", "abnormal relationship" (either word order); or "vein located" followed within 8 characters in the same clause by "left"; or "artery located" followed within 8 characters in the same clause by "right"',
 ('Whirlpool / swirl / spiral appearance', 'All'): '"whirlpool" (two written forms of the Chinese word), "eddy flow", "spiral"; counted only where none of the negation cues "not seen", "not detected", "not shown", "not found", "no obvious", "not obvious", "not heard" occurs in the 12 characters preceding the match within the same clause',
 ('Enteric fluid administered', 'Ultrasound'): '"drinking water"; "oral" followed within 6 characters by "water", "fluid" or "contrast agent"; "water injection"; "injection"; "injection into the stomach"; "warm boiled water"; "via gastric tube"',
 ('Dynamic (real-time) assessment', 'Ultrasound'): '"dynamic"; "real-time observation"; "continuous observation"',
 ('Graded compression', 'Ultrasound'): '"compression"; "probe pressure"; "compression examination"',
 ('Colour Doppler used', 'Ultrasound'): 'CDFI; "color Doppler"; "Doppler"',
 ('Bowel gas explicitly limiting the study', 'Ultrasound'): 'clause containing both a gas term ("bowel gas", "accumulated gas", "gas interference", "gas obscuring", "gas covering", "abundant gas", "strongly reflective gas echoes", "filled with gas") and an interference term ("interference", "obscuring", "covering", "poorly displayed", "incompletely displayed", "unsatisfactory visualization", "unsatisfactory")',
 ('Caecal position', 'Ultrasound / UGI'): '"ileocecal" (cecal region)',
 ('Bedside examination', 'Ultrasound'): '"bedside" (in the examination name)',
 ('Examination type: great-vessel study', 'Ultrasound'): '"abdominal great vessels" (in the examination name)',
 ('Examination type: gastrointestinal ultrasound', 'Ultrasound'): '"gastrointestinal tract" (in the examination name)',
 ('Contrast enhancement', 'CT'): '"enhanced" (in the examination name)',
 ('Three-dimensional reconstruction', 'CT'): '"three-dimensional reconstruction" (in the examination name)',
 ('Duodenojejunal junction', 'UGI'): '"jejunal flexure"; the Chinese short form of the ligament of Treitz; "Treitz"; "suspensory ligament"; "duodenum-jejunum junction" (with an optional hyphen, "with" or "and" between the two words)',
 ('Corkscrew / spring appearance', 'UGI'): '"spring"; "spiral"; "noose-like twist"; "coiled"',
 ('Jejunal position', 'UGI'): '"jejunum" followed within 10 characters by "located in", "situated in" or "deviated"',
 ('Whole-gastrointestinal study', 'UGI'): '"whole gastrointestinal tract" (in the examination name)',
 ('Barium contrast', 'UGI'): '"barium" (in the examination name)',
}
_H = D['H']
assert len(_H) == len(H_EN) + 1 and all((r[0], r[1]) in H_EN for r in _H[1:]), 'Table S1 row without an English translation'
D['H_EN'] = [[_H[0][0], _H[0][1], 'Terms matched (English translation of the Chinese-language pattern)']] + [[r[0], r[1], H_EN[(r[0], r[1])]] for r in _H[1:]]
TAB={'H':('Table S1','Report-content audit patterns, in English translation (original regular expressions in Supplement_1_classifier.py)',D['H_EN']),
     'K':('Table S2','Agreement of the two independent readers on the twelve ultrasound content items, and agreement of the text patterns with their consensus',D['K']),
     'S1':('Table S3','Distribution of final labels and certainty tiers, by modality',D['S1']),
     'G':('Table S4','Between-modality comparison of report-level detection (generalized estimating equation)',D['T4']),
     'P':('Table S5','Report-level detection in the selected subgroup receiving all three examinations, overall and restricted to examinations performed close together in time',D['T5']),
     'S2':('Table S6','Report-level detection stratified by midgut volvulus and by age category',D['S2']),
     'S2B':('Table S7','Detection of a volvulus-specific sign among children with surgically confirmed midgut volvulus',D['S2b']),
     'S3':('Table S18','Documented content of the CT and upper gastrointestinal series index examinations',D['S3']),
     'S4':('Table S8','Ultrasound temporal model with the era boundary placed at 2019, 2020, 2021 and 2022',D['S4']),
     'S5':('Table S9','Ultrasound content audit taking the earliest rather than the closest preoperative examination episode as the index unit',D['S5']),
     'S6':('Table S10','Prevalence of midgut volvulus, and the whirlpool sign among those children, under three definitions of volvulus',D['S6']),
     'S11':('Table S11','Average marginal effect of later era, with bootstrap confidence intervals',D['S11']),
     'S12':('Table S12','Paired differences in detection in the subgroup receiving all three examinations',D['S12']),
     'S13':('Table S13','Interaction terms, and the separated volvulus contrast under penalized likelihood',D['S13']),
     'HP16':('Table S14','Between-modality comparison under four definitions of a positive report',D['S16']),
     'HP17':('Table S15','Standardized detection and differences between modalities under four definitions of a positive report',D['S17']),
     'HP18':('Table S16','Interval from the index examination to operation, choice of index examination, and detection within 2 days of operation',D['S18']),
     'HP19':('Table S17','Ultrasound content documentation by booking category and in alternative denominators',D['S19']),
     'MRD':('Table S19','Minimum reporting dataset for ultrasound in suspected intestinal malrotation, with the documentation rate in this audit as the baseline for re-audit',D['MRD'])}
def make(src,outfile,figs=None):
    doc=docx.Document(); apply_house_style(doc); add_page_numbers(doc)
    st=doc.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11)
    for s in doc.sections: s.left_margin=s.right_margin=Inches(1.0)
    def para(t,style=None,size=None,italic=False):
        p=doc.add_paragraph(style=style)
        t=ama_p(t)
        for pt in re.split(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)',t):
            if not pt: continue
            if pt.startswith('**') and pt.endswith('**'): r=p.add_run(pt[2:-2]); r.bold=True
            elif pt.startswith('*') and pt.endswith('*') and len(pt)>2: r=p.add_run(pt[1:-1]); r.italic=True
            elif pt.startswith('`') and pt.endswith('`') and len(pt)>2: r=p.add_run(pt[1:-1]); r.font.name='Courier New'
            else: r=p.add_run(pt)
            if italic: r.italic=True
            if size: r.font.size=Pt(size)
        p.paragraph_format.space_after=Pt(8); return p
    def table(key):
        tag,title,data=TAB[key]
        p=para(f'{tag}. {title}'); p.runs[0].bold=True
        make_three_line_table(doc,us_table(data))
        para('')
    for ln in open(src).read().split('\n'):
        ln=ln.strip()
        if not ln: continue
        if ln.startswith('#T '): para(ln[3:],style='Title')
        elif ln.startswith('#H1 '): doc.add_heading(ln[4:],level=1)
        elif ln.startswith('#N '): para(ln[3:])
        elif ln.startswith('#FIG'):
            doc.add_picture(OUT+'FigS1_paired_subgroup.png',width=Inches(6.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif ln.startswith('#TAB'): table(ln[4:])
        else: para(ln)
    _blank_props(doc).save(OUT+outfile); print('saved',outfile)
make('supp1.md','EJR_Supplement_1_NLP_and_report_audit.docx')
make('supp2.md','EJR_Supplement_2_models_and_subgroups.docx')
make('supp3.md','EJR_Supplement_3_CT_and_UGI_content_audit.docx')
make('supp4.md','EJR_Supplement_4_minimum_reporting_dataset.docx')
# the runnable rule set deposited with Supplement 1 is the analysis copy, unchanged
shutil.copyfile(AN+'classifier.py',OUT+'Supplement_1_classifier.py')
with open(OUT+'Supplement_1_classifier.py','a',encoding='utf-8') as f:
    f.write("\n\n# " + "-"*75 + "\n# Original-language material referred to in Supplement 1. The study reports are written in Chinese,\n# so these strings are Chinese by necessity; Supplement 1 gives their English translation.\n# " + "-"*75 + "\n")
    f.write("# Search of the institutional research database for operative records (Supplement 1, Section C): field name -> string,\n# matched case-insensitively. Field names: operative diagnosis, procedure name.\n")
    f.write("OPERATIVE_SEARCH = {'术中诊断': '肠旋转不良', '手术名称': 'Ladd'}\n\n")
    f.write("# Report-content audit patterns (Table S1 of Supplement 1): (element, modality, pattern as used, verbatim).\n# Where a row holds a note in English after the pattern, the note describes a further condition applied in the audit.\nCONTENT_PATTERNS = [\n")
    for r in D['H'][1:]:
        f.write('    (%r, %r, %r),\n' % (us(r[0]), r[1], r[2]))
    f.write("]\n")
print('saved Supplement_1_classifier.py (with OPERATIVE_SEARCH and CONTENT_PATTERNS)')

# clean the file properties of every package document (author, times, thumbnail, embedded-image text chunks)
from scrub_metadata import scrub_all
scrub_all()
